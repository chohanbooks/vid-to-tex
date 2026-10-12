import os
import tempfile
import streamlit as st
from groq import Groq
import yt_dlp
from docx import Document

# Initialize Groq client using Streamlit secrets
client = Groq(api_key=st.secrets["GROQ_API_KEY"])

st.set_page_config(page_title="Groq YouTube Bilingual Transcriber", page_icon="⚡", layout="centered")

st.title("⚡ YouTube Transcriber (Groq + Urdu/English Mix)")
st.markdown("Blazing-fast transcription for YouTube videos containing conversational Urdu and English mixing, outputting a clean Word document (`.docx`).")

# --- PIPELINE STAGE 1: INPUT & CONFIGURATION ---
youtube_url = st.text_input("YouTube Video URL", placeholder="https://www.youtube.com/watch?v=...")

language_option = st.selectbox(
    "Select Language Mode",
    options=["Auto-Detect (Best for Mixed Urdu/English)", "Urdu (ur)", "English (en)"]
)

lang_code = None
if "Urdu (ur)" in language_option:
    lang_code = "ur"
elif "English (en)" in language_option:
    lang_code = "en"


# --- PIPELINE STAGE 2: AUDIO EXTRACTION ---
def download_audio(url: str, output_dir: str) -> tuple:
    """Downloads YouTube audio stream and returns file path and title."""
    output_template = os.path.join(output_dir, "audio.%(ext)s")
    ydl_opts = {
        'format': 'bestaudio/best',
        'postprocessors': [{
            'key': 'FFmpegExtractAudio',
            'preferredcodec': 'mp3',
            'preferredquality': '192',
        }],
        'outtmpl': output_template,
        'quiet': True
    }
    with yt_dlp.YoutubeDL(ydl_opts) as ydl:
        info = ydl.extract_info(url, download=True)
        video_title = info.get('title', 'video_transcript')
        audio_path = os.path.join(output_dir, "audio.mp3")
    return audio_path, video_title


# --- PIPELINE STAGE 3: AI PROCESSING (GROQ WHISPER) ---
def transcribe_with_groq(audio_path: str, language: str = None) -> str:
    """Sends audio file to Groq's ultra-fast whisper-large-v3 model."""
    with open(audio_path, "rb") as audio_file:
        kwargs = {
            "model": "whisper-large-v3",
            "file": audio_file,
            "response_format": "text"
        }
        if language:
            kwargs["language"] = language
            
        transcript = client.audio.transcriptions.create(**kwargs)
    return transcript


# --- PIPELINE STAGE 4: OUTPUT DOCUMENT GENERATION ---
def create_word_document(title: str, url: str, text: str, output_dir: str) -> str:
    """Builds a formatted Word document (.docx) from the transcript."""
    doc = Document()
    doc.add_heading(f"Transcription: {title}", level=1)
    doc.add_paragraph(f"Source URL: {url}\n")
    doc.add_heading("Transcript Text", level=2)
    doc.add_paragraph(text)
    
    doc_path = os.path.join(output_dir, "transcript.docx")
    doc.save(doc_path)
    return doc_path


# --- EXECUTION PIPELINE TRIGGER ---
if st.button("Run Groq Transcription Pipeline", type="primary"):
    if not youtube_url:
        st.warning("Please enter a valid YouTube URL first.")
    else:
        temp_dir = tempfile.mkdtemp()
        
        try:
            with st.spinner("Step 2/4: Extracting audio stream from YouTube..."):
                audio_path, video_title = download_audio(youtube_url, temp_dir)
                
            with st.spinner("Step 3/4: Processing via Groq Whisper API (Urdu/English code-switching)..."):
                transcript_text = transcribe_with_groq(audio_path, lang_code)
                
            with st.spinner("Step 4/4: Formatting and compiling Word document..."):
                docx_file_path = create_word_document(video_title, youtube_url, transcript_text, temp_dir)
                
            st.success("Pipeline execution complete!")
            
            with st.expander("View Transcript Preview"):
                st.write(transcript_text)
                
            with open(docx_file_path, "rb") as f:
                st.download_button(
                    label="📥 Download Word File (.docx)",
                    data=f,
                    file_name="groq_youtube_transcript.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
                
        except Exception as e:
            st.error(f"Pipeline error encountered: {e}")